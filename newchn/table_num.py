'''
Author:WangPeng  Email:wangpdm@163.com
Date:2019-07-5--15:16
File：table_num.py
Describe:match the chinese-num and alab-num
'''
from __future__ import unicode_literals
import re
import os
import docx
from glob import glob
import json



def write_result(doc, para_no, alab_chn_list, crt_chn_list, cor_ornot):
    val_table_all = []
    doc = os.path.splitext(os.path.split(doc)[1])[0]
    with open('./test_result/alab_chn_cpy-1.txt', 'a') as f:
        dict = zip(para_no, alab_chn_list, crt_chn_list, cor_ornot)

        for fn, ft, m, rw in dict:
            ft_str = ''.join(str(ft).replace('[','',1).rsplit(']',1))
            m_str = ''.join(str(m).replace('[','',1).rsplit(']',1))
            #val_table = json.dumps(doc + ' | ' + fn + ' | ' + ft_str + ' | ' + rw + ' | ' + m_str + '\n')
            #val_table_all.append(val_table)
            f.write(doc + ' | ' + fn + ' | ' + ft_str + ' | ' + rw + ' | ' + m_str + '\n')
    #return val_table_all


# main
def table_num_match(doc_path):
    for doc in sorted(doc_path):
        file = docx.Document(doc)
        tables = file.tables

        print('The Doc name is {}'.format(doc))
        order = 1
        error_para_list = []
        doc_error_list = []
        doc_corr_chn = []
        RoW = []
        error_tab_list = []
        tab_error_list = []
        tab_corr_chn = []
        tRoW = []
        for t in tables:
            #print(t)

            tab_index = 'table'+str(order)
            order += 1
            result = ''
            for i in range(1, len(t.rows)):
                #for j in range(1, len(t.columns)):
                    result += t.cell(i,1).text
         # read the content from word documents
            numlist = {}
            para_index = []
            para_error_list = []
            para_corr_chn = []
            alab_num_1 = []
            chn_num_1 = []
            alab_num_tf = []
            alab_index = []
            chn_index = []
            alab_num = []
            chn_num = []
            #if re.match('\d', para.text):
                #para_num = para.text

            white_text = re.sub("[{}+' '+\s+\!\/_,$%^*(+\"\']+|[【】+——！，。？、~@#￥%……&*（）+]", '', result)
            print('the content: {}'.format(white_text))

            alab_num_iter = re.finditer(r'([\d.]+[万元])', white_text)

            for item in alab_num_iter:
                #print(item.span()[0], item.group())
                alab_index.append(item.span()[0])
                alab_num.append(item.group())

            #print(alab_index)
            chn_num_iter = re.finditer(r'(?<!\d)[壹贰叁肆伍陆柒捌玖万仟千佰拾亿零]+元', white_text)
            #print(len(chn_num))
            for item in chn_num_iter:
                chn_index.append(item.span()[0])
                chn_num.append(item.group())
            #if para.text == '' | '\d':
                #break
            #print("段落数:{}".format(para_index))
            #para_index = para.text[:10]

            #print(chn_index)
            alab_dict = dict(zip(alab_index, alab_num))
            #print(alab_dict)
            chn_dict = dict(zip(chn_index, chn_num))
            #print(chn_dict)
            mth_item = 0
            alab_num_reg = []
            chn_num_reg = []
            if len(alab_num)!=0 and len(chn_num)!=0:

                if len(alab_num) != len(chn_num):   # find the match pairs of (alab_num, chn_num)
                    for item in chn_index:
                        for item1 in alab_index:
                            if 0 < item - item1 < item - mth_item:
                                mth_item = item1
                        alab_num_reg.append(alab_dict[mth_item])
                        #print('alab_num is {}'.format(alab_num_reg))
                        chn_num_reg.append(chn_dict[item])
                        #print('chn_num is {}'.format(chn_num_reg))
                else:
                    #
                    alab_num_reg = alab_num
                    #print('alab_num_reg is {}'.format(alab_num_reg))
                    chn_num_reg = chn_num
                    #print('chn_num_reg is {}'.format(chn_num_reg))
            #else:
                #alab_num_reg = []
                #chn_num_reg = []

            for item in alab_num_reg:
                item = item.strip('元')
                alab_num_1.append(item)  # alab_num_1 去掉‘元’字的序列
            for item in chn_num_reg:
                item = item.strip('元')
                chn_num_1.append(item) # chn_num_1 去掉‘元’字的序列


            #transfer alab_num to chn_num
            for num in alab_num_1:
                if num.find('万'):
                    num_k = num
                    num = num.strip('万')
                length = len(num)
                dict_len ={1:'', 2:'拾', 3:'佰', 4:'仟', 5:'万', 6:'拾万', 7:'佰万', 8:'仟万', 9:'亿'}
                dict_num = {1:'壹', 2:'贰', 3:'叁', 4:'肆', 5:'伍', 6:'陆', 7:'柒', 8:'捌', 9:'玖', 0:'零'}
                num_1 = enumerate(num)
                index=0
                for i, n in num_1:
                    #print(i,n)
                    if n != '0':
                        numlist[index] = [length-i]
                        numlist[index].append(dict_num[int(n)])
                        numlist[index].append(dict_len[length-i])
                        index = index + 1

                for idx in numlist:
                    if idx + 1 >= len(numlist):
                        break
                    elif numlist[idx][0]-numlist[idx+1][0] > 1:
                        numlist[idx].append('零')

                #print(numlist)
                chn_str = ''
                for i in numlist:
                    chn_str = chn_str + ''.join(numlist[i][1:])

                #print(chn_str)
                #print('num_k is {}'.format(num_k))
                if num_k.find('万') != -1:
                    chn_str = chn_str + '万'
                #print('chn_str is {}'.format(chn_str))

                chn_str_1 = re.findall('万',chn_str)

                if len(chn_str_1)>1:
                    #print(len(chn_str_1))
                    chn_str=chn_str.replace('万', '', (chn_str.count('万')-1))

                #print('the alab tranfer to chn is:{}'.format(chn_str))

                alab_num_tf.append(chn_str) # alab数字对应的正确大写数字列表

            #compare chn_str with chn_num
            for index, correct_chn in enumerate(alab_num_tf):
                if correct_chn != chn_num_1[index]:
                    para_error_list.append([alab_num_1[index],chn_num_1[index]])
                    para_corr_chn.append(alab_num_tf[index])

            #print('正确的大写数字是{}'.format(para_corr_chn))
            #if len(para_corr_chn) != 0:

            #print('para_error_list: {}'.format(para_error_list))
            if len(para_error_list) > 0:
                error_para_list.append(tab_index)
                doc_error_list.append(para_error_list)
                doc_corr_chn.append(para_corr_chn)
                RoW.append('不一致')




        #print(error_para_list)
        #print('正确的大写数字是{}'.format(doc_corr_chn))
        #print('******{}'.format(doc_error_list))
        if len(error_para_list) > 0:
            #RoW = '不一致'
            write_result(doc, error_para_list, doc_error_list, doc_corr_chn, RoW)

        else:
            error_para_list = doc_error_list = doc_corr_chn = [' ']
            RoW = ['一致']
            write_result(doc, error_para_list, doc_error_list, doc_corr_chn, RoW)



            #+table.cell(i,2).text+table.cell(i,3).text+table.cell(i,4).text\
            #+table.cell(i,5).text+table.cell(i,6).text+table.cell(i,7).text
        '''
        print(result)
        alab_index, alab_num, chn_index, chn_num = chn_alab_extr(result)
        print(alab_index, alab_num, chn_index, chn_num)
        error_tab_list, tab_error_list, tab_corr_chn, tRoW = chn_alab_match(alab_index, alab_num, chn_index, chn_num)
        print('error_tab_list:{}\ntab_error_list:{}\ntab_corr_chn:{}\ntRoW:{}'.format(error_tab_list, tab_error_list, tab_corr_chn, tRoW))
        '''
